import hashlib
import io
import json
import os
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile

import requests
from openpyxl import Workbook

from feedback_store import get_supabase_config, normalize_text, supabase_headers


DEFAULT_CATALOG_ADDITIONS = Path(__file__).resolve().parent / "data" / "catalog_additions.json"

SUPABASE_TABLE = "catalog_items"

PRICE_LIST_SOURCE = "price_list"

OPTIONAL_KEYS = (
    "supplier",
    "source_url",
    "source_type",
    "material_type",
    "component_type",
    "thickness_mm",
    "brand_model",
    "keywords",
    "synonyms",
    "description",
    "allowed_use",
    "match_priority",
    "needs_review",
    "added_for_part",
    "notes",
    "created_by",
)

EDITABLE_KEYS = (
    "name",
    "category",
    "unit",
    "price",
    "supplier",
    "source_url",
    "material_type",
    "component_type",
    "thickness_mm",
    "brand_model",
    "keywords",
    "synonyms",
    "description",
    "allowed_use",
    "match_priority",
    "needs_review",
    "notes",
    "active",
    "approved",
    "approved_by",
    "approved_at",
)

COLUMNS = (
    "id",
    "name",
    "category",
    "unit",
    "price",
    *OPTIONAL_KEYS,
    "active",
    "approved",
    "approved_by",
    "approved_at",
    "created_at",
    "updated_at",
)

EXCEL_COLUMNS = (
    ("Name", "name"),
    ("Category", "category"),
    ("Unit", "unit"),
    ("Price (EUR)", "price"),
    ("Material_Type", "material_type"),
    ("Component_Type", "component_type"),
    ("Brand_Model", "brand_model"),
    ("Thickness_mm", "thickness_mm"),
    ("Keywords", "keywords"),
    ("Synonyms", "synonyms"),
    ("Description", "description"),
    ("Allowed_Use", "allowed_use"),
    ("Match_Priority", "match_priority"),
    ("Needs_Review", "needs_review"),
    ("Notes", "notes"),
    ("Supplier", "supplier"),
    ("Source_URL", "source_url"),
    ("Active", "active"),
)

_file_lock = threading.Lock()


def _clean_text(value):
    return " ".join(str(value or "").strip().split())


def _identity(item):
    return (
        normalize_text(item.get("name")),
        normalize_text(item.get("category")),
        normalize_text(item.get("unit")),
    )


def _now():
    return datetime.now(timezone.utc).isoformat()


def _item_id(raw):
    if raw.get("id"):
        return str(raw["id"])
    return "local-" + hashlib.sha1("|".join(_identity(raw)).encode("utf-8")).hexdigest()[:16]


def _optional_value(key, value):
    if key == "needs_review":
        return bool(value)
    if key == "thickness_mm":
        try:
            value = float(value)
        except (TypeError, ValueError):
            return None
        return value if value > 0 else None
    value = _clean_text(value)
    return value or None


def _set_aside_unreadable(path):
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S-%f")
    backup = path.with_name(f"{path.stem}.unreadable-{stamp}.json")
    try:
        os.replace(path, backup)
        print("Unreadable catalog additions file kept as:", backup)
    except Exception as error:
        print("Could not set aside unreadable catalog additions file:", error)


def _read_raw_items(path):
    if not path.exists():
        return []
    try:
        text = path.read_text(encoding="utf-8")
    except OSError as error:
        print("Could not open catalog additions file:", error)
        return None
    if not text.strip():
        return []
    try:
        payload = json.loads(text)
    except ValueError as error:
        print("Catalog additions file is damaged:", error)
        _set_aside_unreadable(path)
        return []
    if not isinstance(payload, list):
        print("Catalog additions file does not contain a list.")
        _set_aside_unreadable(path)
        return []
    return payload


def _validated(item):
    name = _clean_text(item.get("name"))
    category = _clean_text(item.get("category"))
    unit = _clean_text(item.get("unit"))

    try:
        price = float(item.get("price"))
    except (TypeError, ValueError) as exc:
        raise ValueError("Catalog price must be a number.") from exc

    if not name:
        raise ValueError("Catalog item name is required.")
    if not category:
        raise ValueError("Catalog category is required.")
    if not unit:
        raise ValueError("Catalog unit is required.")
    if price <= 0:
        raise ValueError("Catalog price must be greater than zero.")

    return {"name": name, "category": category, "unit": unit, "price": price}


def _clean_item(raw, storage):
    if not isinstance(raw, dict):
        return None
    try:
        required = _validated(raw)
    except ValueError:
        return None

    item = {key: raw.get(key) for key in COLUMNS}
    item.update(required)
    if not item.get("supplier") and raw.get("source"):
        item["supplier"] = raw.get("source")
    if not item.get("notes") and raw.get("market_match"):
        item["notes"] = f'Market match: {raw.get("market_match")}'
    item.update({
        "id": _item_id(raw),
        "active": raw.get("active") is not False,
        "approved": raw.get("approved") is not False,
        "needs_review": bool(raw.get("needs_review")),
        "storage": storage,
        "catalog_addition": raw.get("source_type") != PRICE_LIST_SOURCE,
    })
    return item


def _row(item):
    return {key: item.get(key) for key in COLUMNS if key in item}


def _load_supabase_items(url, key):
    response = requests.get(
        f"{url}/rest/v1/{SUPABASE_TABLE}",
        headers=supabase_headers(key),
        params={"select": "*", "order": "created_at.asc"},
        timeout=15,
    )
    response.raise_for_status()
    return response.json()


def _add_supabase_items(items, url, key):
    rows = [{column: item.get(column) for column in COLUMNS} for item in items]
    for start in range(0, len(rows), 500):
        response = requests.post(
            f"{url}/rest/v1/{SUPABASE_TABLE}",
            headers=supabase_headers(key),
            json=rows[start:start + 500],
            timeout=30,
        )
        response.raise_for_status()


def _patch_supabase_item(item_id, changes, url, key):
    response = requests.patch(
        f"{url}/rest/v1/{SUPABASE_TABLE}",
        headers=supabase_headers(key),
        params={"id": f"eq.{item_id}"},
        json=changes,
        timeout=15,
    )
    response.raise_for_status()


def _write_items(path, items):
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(items, ensure_ascii=False, indent=2)
    with NamedTemporaryFile(
        "w",
        encoding="utf-8",
        dir=str(path.parent),
        delete=False,
        suffix=".tmp",
    ) as handle:
        handle.write(payload)
        temp_path = Path(handle.name)
    os.replace(temp_path, path)


def _change_local_items(path, additions=(), changes_by_id=None, delete_id=None):
    changes_by_id = changes_by_id or {}
    with _file_lock:
        items = _read_raw_items(path)
        if items is None:
            raise RuntimeError("Catalog additions file could not be opened, so it was not changed.")
        found = set()
        kept = []
        for raw in items:
            if not isinstance(raw, dict):
                kept.append(raw)
                continue
            item_id = _item_id(raw)
            if item_id == delete_id:
                found.add(item_id)
                continue
            if item_id in changes_by_id:
                found.add(item_id)
                raw = {**raw, "id": item_id, **changes_by_id[item_id]}
            kept.append(raw)
        missing = (set(changes_by_id) | ({delete_id} if delete_id else set())) - found
        if missing:
            raise ValueError("This catalog item no longer exists.")
        kept.extend(additions)
        _write_items(path, kept)


def _supabase_or_error():
    url, key = get_supabase_config()
    if not (url and key):
        raise RuntimeError("Supabase is not configured.")
    return url, key


def list_catalog_items(path=DEFAULT_CATALOG_ADDITIONS, raise_errors=False):
    items = []

    for raw in _read_raw_items(Path(path)) or []:
        item = _clean_item(raw, "local file")
        if item is not None:
            items.append(item)

    url, key = get_supabase_config()
    if url and key:
        try:
            for raw in _load_supabase_items(url, key):
                item = _clean_item(raw, "Supabase")
                if item is not None:
                    items.append(item)
        except Exception as error:
            print("Supabase catalog load failed:", error)
            if raise_errors:
                raise

    return items


def load_catalog_additions(path=DEFAULT_CATALOG_ADDITIONS):
    items = {}
    for item in list_catalog_items(path):
        if item["active"] and item["approved"]:
            items[_identity(item)] = item
    return list(items.values())


def price_list_imported(items):
    return any(item.get("source_type") == PRICE_LIST_SOURCE for item in items)


def load_live_catalog(base_items, path=DEFAULT_CATALOG_ADDITIONS):
    stored = list_catalog_items(path)
    usable = [item for item in stored if item["active"] and item["approved"]]

    if price_list_imported(stored):
        return usable

    items = list(base_items)
    for item in usable:
        if not catalog_item_exists(items, item):
            items.append(item)
    return items


def catalog_storage_name():
    url, key = get_supabase_config()
    if url and key:
        return "Supabase"
    return "local file"


def catalog_item_status(item):
    if not item.get("active", True):
        return "disabled"
    if not item.get("approved", True):
        return "pending"
    return "approved"


def catalog_item_exists(existing_items, candidate):
    candidate_identity = _identity(candidate)
    return any(_identity(item) == candidate_identity for item in existing_items)


def _new_item(item, approved, approved_by=None):
    created_at = _now()
    clean_item = {
        "id": uuid.uuid4().hex,
        **_validated(item),
        "active": item.get("active") is not False,
        "approved": approved,
        "created_at": created_at,
        "updated_at": created_at,
    }
    if approved:
        clean_item["approved_by"] = _clean_text(approved_by) or None
        clean_item["approved_at"] = created_at

    for optional_key in OPTIONAL_KEYS:
        value = _optional_value(optional_key, item.get(optional_key))
        if value not in (None, False):
            clean_item[optional_key] = value

    return clean_item


def add_catalog_item(item, path=DEFAULT_CATALOG_ADDITIONS, approved=False, approved_by=None):
    path = Path(path)

    if approved and not _clean_text(approved_by):
        raise ValueError("Enter your name before adding an approved item.")

    clean_item = _new_item(item, approved, approved_by)
    identity = _identity(clean_item)

    for existing in list_catalog_items(path, raise_errors=approved):
        if _identity(existing) == identity:
            return False, existing

    url, key = get_supabase_config()
    if url and key:
        try:
            _add_supabase_items([clean_item], url, key)
            print("Added catalog item to Supabase:", clean_item["name"])
            return True, _clean_item(clean_item, "Supabase")
        except Exception as error:
            if approved:
                raise
            print("Supabase catalog save failed, saving locally:", error)

    with _file_lock:
        items = _read_raw_items(path)
        if items is None:
            raise RuntimeError("Catalog additions file could not be opened, so it was not overwritten.")
        for existing in items:
            if isinstance(existing, dict) and _identity(existing) == identity:
                return False, _clean_item(existing, "local file") or existing
        items.append(clean_item)
        _write_items(path, items)

    print("Added catalog item locally:", clean_item["name"])
    return True, _clean_item(clean_item, "local file")


def _apply_changes(item, changes, path):
    if item.get("storage") == "Supabase":
        url, key = _supabase_or_error()
        _patch_supabase_item(item["id"], changes, url, key)
    else:
        _change_local_items(path, changes_by_id={item["id"]: changes})


def _clean_changes(item, changes):
    changes = {key: value for key, value in changes.items() if key in EDITABLE_KEYS}
    merged = {**item, **changes}
    changes.update(_validated(merged))
    for optional_key in OPTIONAL_KEYS:
        if optional_key in changes:
            changes[optional_key] = _optional_value(optional_key, changes[optional_key])
    for flag in ("active", "approved"):
        if flag in changes:
            changes[flag] = bool(changes[flag])
    return changes


def update_catalog_item(item, changes, path=DEFAULT_CATALOG_ADDITIONS):
    path = Path(path)
    changes = _clean_changes(item, changes)
    changes["updated_at"] = _now()

    identity = _identity({**item, **changes})
    for other in list_catalog_items(path, raise_errors=True):
        if other["id"] != item["id"] and _identity(other) == identity:
            raise ValueError("Another catalog item already has this name, category and unit.")

    _apply_changes(item, changes, path)
    return {**item, **changes}


def approve_catalog_item(item, approved_by, changes=None, path=DEFAULT_CATALOG_ADDITIONS):
    approved_by = _clean_text(approved_by)
    if not approved_by:
        raise ValueError("Enter your name before approving.")
    return update_catalog_item(
        item,
        {**(changes or {}), "approved": True, "approved_by": approved_by, "approved_at": _now()},
        path,
    )


def delete_catalog_item(item, path=DEFAULT_CATALOG_ADDITIONS):
    if item.get("storage") == "Supabase":
        url, key = _supabase_or_error()
        response = requests.delete(
            f"{url}/rest/v1/{SUPABASE_TABLE}",
            headers=supabase_headers(key),
            params={"id": f'eq.{item["id"]}'},
            timeout=15,
        )
        response.raise_for_status()
    else:
        _change_local_items(Path(path), delete_id=item["id"])


def import_catalog_items(
    rows,
    approved_by,
    path=DEFAULT_CATALOG_ADDITIONS,
    source_type=PRICE_LIST_SOURCE,
    dry_run=False,
):
    path = Path(path)
    approved_by = _clean_text(approved_by)
    if not approved_by and not dry_run:
        raise ValueError("Enter your name before importing.")

    existing = {}
    for item in list_catalog_items(path, raise_errors=True):
        existing.setdefault(_identity(item), item)

    result = {"added": [], "updated": [], "unchanged": 0, "skipped": [], "duplicates": []}
    new_items = []
    updates = []
    seen = set()

    for row in rows:
        try:
            required = _validated(row)
        except ValueError as error:
            result["skipped"].append(f'{row.get("name") or "(no name)"}: {error}')
            continue

        identity = _identity(required)
        if identity in seen:
            result["duplicates"].append(required["name"])
            continue
        seen.add(identity)

        current = existing.get(identity)
        if current is None:
            new_items.append(_new_item(
                {"source_type": source_type, **row},
                approved=True,
                approved_by=approved_by,
            ))
            result["added"].append(required["name"])
            continue

        fields = dict(required)
        for key in OPTIONAL_KEYS:
            if key in row and key not in ("source_type", "added_for_part", "created_by"):
                fields[key] = _optional_value(key, row[key])
        if "active" in row:
            fields["active"] = bool(row["active"])
        if not current["approved"]:
            fields.update({"approved": True, "approved_by": approved_by, "approved_at": _now()})

        changes = {key: value for key, value in fields.items() if current.get(key) != value}
        if changes:
            changes["updated_at"] = _now()
            updates.append((current, changes))
            result["updated"].append(required["name"])
        else:
            result["unchanged"] += 1

    if dry_run:
        return result

    url, key = get_supabase_config()
    if url and key:
        _add_supabase_items(new_items, url, key)
        for current, changes in updates:
            _apply_changes(current, changes, path)
    else:
        _change_local_items(
            path,
            additions=new_items,
            changes_by_id={current["id"]: changes for current, changes in updates},
        )

    return result


def catalog_to_excel(items, markups, project_multiplier):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "All_Items"
    sheet.append([header for header, _ in EXCEL_COLUMNS])

    for item in items:
        row = []
        for _, key in EXCEL_COLUMNS:
            value = item.get(key)
            if key in ("needs_review", "active"):
                value = "YES" if (value if key == "needs_review" else value is not False) else "NO"
            row.append(value)
        sheet.append(row)

    categories = workbook.create_sheet("Categories")
    categories.append(["Category tag (used in project line items)", "Markup coefficient"])
    for category, coefficient in markups.items():
        categories.append([f"{category} ({category.upper()})", coefficient])

    settings = workbook.create_sheet("Global_Settings")
    settings.append(["Setting", "Value", "Note"])
    settings.append(["Overhead / margin multiplier", project_multiplier, ""])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
