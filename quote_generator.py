from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_SUPPLIER_NAME = 'IĮ "Buki kaltai"'
DEFAULT_SUPPLIER_DETAILS = [
    "305189636",
    "LT100012484610",
    "Šilo g. 6, Lekėčiai, Šakių r.",
    "062533752",
]


def safe_number(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return float(default)


def sanitize_filename(value):
    value = str(value or "quote").strip()
    cleaned = "".join(
        char if char.isalnum() or char in ("-", "_") else "_"
        for char in value
    )
    cleaned = cleaned.strip("_")
    return cleaned or "quote"


def calculate_quote_totals(
    furniture_price,
    services,
    vat_percent=21.0,
    discount_value=0.0,
    discount_mode="amount",
):
    furniture_price = max(0.0, safe_number(furniture_price))

    clean_services = []
    for service in services or []:
        name = str(service.get("name") or "").strip()
        price = max(0.0, safe_number(service.get("price")))
        quantity = max(0.0, safe_number(service.get("quantity"), 1.0))

        if not name or price <= 0 or quantity <= 0:
            continue

        clean_services.append({
            "name": name,
            "price": price,
            "quantity": quantity,
            "total": price * quantity,
        })

    subtotal_before_discount = furniture_price + sum(
        service["total"] for service in clean_services
    )

    discount_value = max(0.0, safe_number(discount_value))

    if discount_mode == "percent":
        discount_amount = subtotal_before_discount * min(discount_value, 100.0) / 100.0
    else:
        discount_amount = min(discount_value, subtotal_before_discount)

    subtotal_after_discount = max(0.0, subtotal_before_discount - discount_amount)
    vat_percent = max(0.0, safe_number(vat_percent))
    vat_amount = subtotal_after_discount * vat_percent / 100.0
    total_with_vat = subtotal_after_discount + vat_amount

    return {
        "furniture_price": furniture_price,
        "services": clean_services,
        "subtotal_before_discount": subtotal_before_discount,
        "discount_amount": discount_amount,
        "subtotal_after_discount": subtotal_after_discount,
        "vat_percent": vat_percent,
        "vat_amount": vat_amount,
        "total_with_vat": total_with_vat,
    }


def build_quote_excel(
    quote_number,
    quote_date,
    validity_date,
    client_name,
    furniture_name,
    furniture_price,
    services,
    vat_percent=21.0,
    discount_value=0.0,
    discount_mode="amount",
    supplier_name=DEFAULT_SUPPLIER_NAME,
    supplier_details=None,
    notes=None,
    logo_path=Path("data") / "logo.jpg",
):
    totals = calculate_quote_totals(
        furniture_price=furniture_price,
        services=services,
        vat_percent=vat_percent,
        discount_value=discount_value,
        discount_mode=discount_mode,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Komerciinis pasiulymas"
    ws.sheet_view.showGridLines = False

    # Page/layout
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "A15"
    ws.print_area = "A1:G40"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

    widths = {
        "A": 6,
        "B": 31,
        "C": 14,
        "D": 10,
        "E": 16,
        "F": 14,
        "G": 20,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    title_font = Font(name="Times New Roman", size=16, bold=True)
    header_font = Font(name="Times New Roman", size=10, bold=True)
    normal_font = Font(name="Times New Roman", size=10)
    total_font = Font(name="Times New Roman", size=11, bold=True)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="EAD5DF")

    # Logo + title
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 18

    logo_file = Path(logo_path) if logo_path else None
    if logo_file and logo_file.exists():
        try:
            logo = XLImage(str(logo_file))
            # Keep the image compact so the quote still prints cleanly on one page.
            max_width = 150
            max_height = 65
            scale = min(
                max_width / float(logo.width or max_width),
                max_height / float(logo.height or max_height),
                1.0,
            )
            logo.width = int(logo.width * scale)
            logo.height = int(logo.height * scale)
            ws.add_image(logo, "A1")
        except Exception as error:
            print("Could not add quote logo:", error)

    ws.merge_cells("C2:G2")
    ws["C2"] = "KOMERCINIS PASIŪLYMAS"
    ws["C2"].font = title_font
    ws["C2"].alignment = Alignment(horizontal="center", vertical="center")

    # Quote metadata
    metadata = [
        ("Numeris", quote_number),
        ("Data", quote_date),
        ("Galiojimo data", validity_date),
    ]
    start_row = 5
    for offset, (label, value) in enumerate(metadata):
        row = start_row + offset
        ws[f"B{row}"] = label
        ws[f"B{row}"].font = header_font
        ws[f"C{row}"] = value
        ws[f"C{row}"].font = normal_font
        if hasattr(value, "strftime"):
            ws[f"C{row}"].number_format = "dd/mm/yyyy"

    # Supplier/customer blocks
    ws["B9"] = "Tiekėjas"
    ws["B9"].font = header_font
    ws["B10"] = supplier_name
    ws["B10"].font = normal_font

    supplier_details = supplier_details or []
    for index, detail in enumerate(supplier_details[:4], start=11):
        ws[f"B{index}"] = str(detail)
        ws[f"B{index}"].font = normal_font

    ws["E9"] = "Užsakovas"
    ws["E9"].font = header_font
    ws["E10"] = client_name
    ws["E10"].font = normal_font

    # Table
    table_row = 15
    headers = [
        "Nr.",
        "Produktas/paslauga",
        "Kaina",
        "Kiekis",
        "Kaina už visus",
        "PVM",
        "Kaina su mokesčiais",
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=table_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = [{
        "name": str(furniture_name or "Furniture"),
        "price": totals["furniture_price"],
        "quantity": 1.0,
        "total": totals["furniture_price"],
    }]
    rows.extend(totals["services"])

    current_row = table_row + 1
    for number, item in enumerate(rows, start=1):
        line_total = safe_number(item.get("total"))
        line_vat = line_total * totals["vat_percent"] / 100.0
        line_gross = line_total + line_vat

        values = [
            number,
            item.get("name"),
            safe_number(item.get("price")),
            safe_number(item.get("quantity"), 1.0),
            line_total,
            line_vat,
            line_gross,
        ]

        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=current_row, column=col, value=value)
            cell.font = normal_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center" if col != 2 else "left",
                vertical="center",
                wrap_text=True,
            )
            if col in (3, 5, 6, 7):
                cell.number_format = '#,##0.00 "€"'
            elif col == 4:
                cell.number_format = "0.##"

        current_row += 1

    # Summary
    summary_row = current_row + 1
    summary = [
        ("Kaina", totals["subtotal_before_discount"]),
    ]

    if totals["discount_amount"] > 0:
        summary.append(("Nuolaida", -totals["discount_amount"]))

    summary.extend([
        (f'PVM {totals["vat_percent"]:g}%', totals["vat_amount"]),
        ("Viso", totals["total_with_vat"]),
    ])

    for index, (label, value) in enumerate(summary):
        row = summary_row + index
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        label_cell = ws.cell(row=row, column=5, value=label)
        value_cell = ws.cell(row=row, column=7, value=value)
        label_cell.alignment = Alignment(horizontal="right")
        value_cell.alignment = Alignment(horizontal="right")
        value_cell.number_format = '#,##0.00 "€"'

        if label == "Viso":
            label_cell.font = total_font
            value_cell.font = total_font
            label_cell.border = Border(top=thin)
            value_cell.border = Border(top=thin)
        else:
            label_cell.font = header_font
            value_cell.font = header_font

    # Notes
    note_row = summary_row + len(summary) + 3
    default_notes = [
        "Kainos gali keistis pasikeitus žaliavų ar paslaugų kainoms.",
        "Pasiūlymas galioja iki nurodytos galiojimo datos.",
    ]
    note_lines = notes if notes is not None else default_notes

    for note in note_lines:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
        cell = ws.cell(row=note_row, column=1, value=f"• {note}")
        cell.font = Font(name="Times New Roman", size=9)
        cell.alignment = Alignment(wrap_text=True)
        note_row += 1

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def quote_filename(quote_number, client_name):
    quote_part = sanitize_filename(quote_number)
    client_part = sanitize_filename(client_name)
    return f"{quote_part}_{client_part}.xlsx"
